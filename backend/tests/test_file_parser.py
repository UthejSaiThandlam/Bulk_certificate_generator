import io
import json
import pytest
import openpyxl
import docx
from app.utils.file_parser import extract_recipients_from_any_file


def test_parse_csv_file():
    csv_bytes = b"name,email,course,date\nAlice Cooper,alice@test.com,Cybersecurity,2025-05-01\n"
    recipients = extract_recipients_from_any_file("test.csv", csv_bytes)
    assert len(recipients) == 1
    assert recipients[0].recipient_name == "Alice Cooper"
    assert recipients[0].recipient_email == "alice@test.com"
    assert recipients[0].course_name == "Cybersecurity"


def test_parse_json_file():
    json_data = [
        {"recipient_name": "Bob Marley", "recipient_email": "bob@reggae.com", "course_name": "Music Theory"}
    ]
    json_bytes = json.dumps(json_data).encode("utf-8")
    recipients = extract_recipients_from_any_file("students.json", json_bytes)
    assert len(recipients) == 1
    assert recipients[0].recipient_name == "Bob Marley"


def test_parse_excel_file():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Student Name", "Contact Email", "Program", "Award Date"])
    ws.append(["Clara Schumann", "clara@piano.com", "Classical Composition", "2025-04-12"])

    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_bytes = excel_buffer.getvalue()

    recipients = extract_recipients_from_any_file("roster.xlsx", excel_bytes)
    assert len(recipients) == 1
    assert recipients[0].recipient_name == "Clara Schumann"
    assert recipients[0].recipient_email == "clara@piano.com"
    assert recipients[0].course_name == "Classical Composition"


def test_parse_docx_file():
    doc = docx.Document()
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Full Name"
    hdr_cells[1].text = "Email Address"
    hdr_cells[2].text = "Certification"

    row_cells = table.add_row().cells
    row_cells[0].text = "David Bowie"
    row_cells[1].text = "david@space.com"
    row_cells[2].text = "Glam Rock History"

    doc_buffer = io.BytesIO()
    doc.save(doc_buffer)
    doc_bytes = doc_buffer.getvalue()

    recipients = extract_recipients_from_any_file("participants.docx", doc_bytes)
    assert len(recipients) == 1
    assert recipients[0].recipient_name == "David Bowie"
    assert recipients[0].recipient_email == "david@space.com"
    assert recipients[0].course_name == "Glam Rock History"


def test_upload_excel_via_api(client):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["name", "email", "course"])
    ws.append(["Ella Fitzgerald", "ella@jazz.com", "Vocal Performance"])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    files = {"file": ("roster.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    response = client.post("/api/certificates/upload-file", files=files)
    assert response.status_code == 202
    assert response.json()["total"] == 1


def test_upload_docx_via_api(client):
    doc = docx.Document()
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "name"
    table.rows[0].cells[1].text = "email"
    table.rows[0].cells[2].text = "course"
    row = table.add_row()
    row.cells[0].text = "Frank Sinatra"
    row.cells[1].text = "frank@swing.com"
    row.cells[2].text = "Big Band Jazz"

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    files = {"file": ("students.docx", buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    response = client.post("/api/certificates/upload-file", files=files)
    assert response.status_code == 202
    assert response.json()["total"] == 1
